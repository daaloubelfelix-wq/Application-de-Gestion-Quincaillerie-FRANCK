"""
Export de données vers des fichiers Excel (.xlsx) — pour partager le
stock avec un fournisseur ou transmettre un rapport de ventes à un
comptable, en dehors de l'application.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_COULEUR_ENTETE = "152C4D"


def _ecrire_entete(feuille, colonnes):
    feuille.append(colonnes)
    for index in range(1, len(colonnes) + 1):
        cellule = feuille.cell(row=1, column=index)
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor=_COULEUR_ENTETE)
        cellule.alignment = Alignment(horizontal="center")


def _ajuster_largeurs_colonnes(feuille):
    for colonne in feuille.columns:
        valeurs = [str(cellule.value) for cellule in colonne if cellule.value is not None]
        largeur = max((len(v) for v in valeurs), default=8)
        feuille.column_dimensions[get_column_letter(colonne[0].column)].width = min(largeur + 3, 40)


def exporter_articles_excel(chemin_fichier, articles, inclure_prix=True):
    """articles : liste de dicts (voir modules.articles.lister_articles).
    inclure_prix=False pour un export sans montants (agent stock, qui ne
    gère que l'inventaire — voir ui/gestion_articles.py)."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Stock"
    colonnes = ["Nom", "Site"] + (["Prix de vente (FCFA)"] if inclure_prix else []) + ["Stock", "Seuil d'alerte"]
    _ecrire_entete(feuille, colonnes)
    for article in articles:
        ligne = [article["nom"], article.get("site_nom") or ""]
        if inclure_prix:
            ligne.append(article["prix_vente"])
        ligne += [article["quantite_stock"], article["seuil_alerte"]]
        feuille.append(ligne)
    _ajuster_largeurs_colonnes(feuille)
    classeur.save(chemin_fichier)


def exporter_rapport_excel(chemin_fichier, periode_texte, totaux, produits):
    """
    totaux : dict avec total_ventes_ttc et nombre_ventes (voir
    modules.rapports.totaux_periode). produits : liste de dicts avec
    nom et quantite_vendue (voir modules.rapports.produits_plus_vendus).
    """
    classeur = Workbook()

    feuille_resume = classeur.active
    feuille_resume.title = "Résumé"
    feuille_resume.append(["Période", periode_texte])
    feuille_resume.append(["Total ventes TTC (FCFA)", totaux["total_ventes_ttc"]])
    feuille_resume.append(["Nombre de ventes", totaux["nombre_ventes"]])
    for ligne in feuille_resume.iter_rows(min_row=1, max_row=3, min_col=1, max_col=1):
        ligne[0].font = Font(bold=True)
    feuille_resume.column_dimensions["A"].width = 28
    feuille_resume.column_dimensions["B"].width = 22

    feuille_produits = classeur.create_sheet("Produits les plus vendus")
    _ecrire_entete(feuille_produits, ["Article", "Quantité vendue"])
    for produit in produits:
        feuille_produits.append([produit["nom"], produit["quantite_vendue"]])
    _ajuster_largeurs_colonnes(feuille_produits)

    classeur.save(chemin_fichier)
